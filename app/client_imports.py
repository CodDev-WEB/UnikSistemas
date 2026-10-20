import csv
import io
import re
import unicodedata
from collections.abc import Iterable
from pathlib import PurePosixPath
from zipfile import BadZipFile

from openpyxl import load_workbook
from pydantic import ValidationError

from app.schemas import ClientCreate

MAX_IMPORT_ROWS = 1000

FIELD_ALIASES = {
    "legal_name": {"legalname", "razaosocial", "nomeempresarial", "empresa", "nome"},
    "trade_name": {"tradename", "nomefantasia", "fantasia"},
    "tax_id": {"taxid", "cnpj", "cpf", "cnpjcpf", "documento", "identificadorfiscal"},
    "email": {"email", "emailcontato", "correioeletronico"},
    "phone": {"phone", "telefone", "celular", "whatsapp"},
    "status": {"status", "situacao"},
}
FIELD_LABELS = {
    "legal_name": "Razão social",
    "trade_name": "Nome fantasia",
    "tax_id": "CNPJ/CPF",
    "email": "E-mail",
    "phone": "Telefone",
    "status": "Status",
}


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]", "", text.lower())


def canonical_tax_id(value: str | None) -> str | None:
    if not value or not value.strip():
        return None
    cleaned = value.strip()
    if re.fullmatch(r"[\d.\-/\s]+", cleaned):
        digits = re.sub(r"\D", "", cleaned)
        return digits or None
    return cleaned.casefold()


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _records_from_csv(content: bytes) -> tuple[list[object], Iterable[list[object]]]:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            decoded = content.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise ValueError("Não foi possível ler o CSV. Salve o arquivo em UTF-8 ou Excel (.xlsx).") from exc
    sample = decoded[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;|\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(decoded, newline=""), dialect)
    all_rows = iter(reader)
    try:
        headers = next(all_rows)
    except StopIteration:
        raise ValueError("O arquivo está vazio.") from None
    return headers, all_rows


def _records_from_xlsx(content: bytes) -> tuple[list[object], Iterable[list[object]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        headers = list(next(iterator, ()))
        rows = (list(row) for row in iterator)
        return headers, rows
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise ValueError("Arquivo Excel inválido ou danificado.") from exc


def parse_client_import(filename: str, content: bytes) -> list[dict]:
    extension = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise ValueError("Formato não suportado. Envie um arquivo .csv ou .xlsx.")
    headers, rows = (
        _records_from_csv(content) if extension == ".csv" else _records_from_xlsx(content)
    )

    columns: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalized = normalize_header(header)
        for field, aliases in FIELD_ALIASES.items():
            if normalized in aliases and field not in columns:
                columns[field] = index
                break
    if "legal_name" not in columns:
        raise ValueError(
            "Não encontrei a coluna obrigatória de razão social. Use 'Razão Social' ou 'legal_name'."
        )

    parsed: list[dict] = []
    for row_number, cells in enumerate(rows, start=2):
        values = [_cell_text(value) for value in cells]
        if not any(value.strip() for value in values):
            continue
        if len(parsed) >= MAX_IMPORT_ROWS:
            raise ValueError(f"O arquivo excede o limite de {MAX_IMPORT_ROWS} linhas de dados.")
        row = {"row_number": row_number}
        for field, index in columns.items():
            row[field] = values[index] if index < len(values) and values[index].strip() else None
        parsed.append(row)

    if not parsed:
        raise ValueError("Não há linhas de dados para importar.")
    return parsed


def validate_import_rows(rows: list[dict], existing_tax_ids: set[str]) -> list[dict]:
    preview: list[dict] = []
    seen_tax_ids: set[str] = set()
    for source in rows:
        row_number = source["row_number"]
        raw_status = (source.get("status") or "").strip().casefold()
        status_map = {
            "": "active",
            "active": "active",
            "ativo": "active",
            "ativa": "active",
            "archived": "archived",
            "arquivado": "archived",
            "arquivada": "archived",
        }
        values = {
            "legal_name": (source.get("legal_name") or "").strip(),
            "trade_name": (source.get("trade_name") or "").strip() or None,
            "tax_id": (source.get("tax_id") or "").strip() or None,
            "email": (source.get("email") or "").strip() or None,
            "phone": (source.get("phone") or "").strip() or None,
            "status": status_map.get(raw_status, "active"),
        }
        errors: list[str] = []
        if raw_status not in status_map:
            errors.append("Status inválido; use Ativo ou Arquivado.")
        try:
            validated = ClientCreate.model_validate(values)
            values = validated.model_dump(mode="json")
        except ValidationError as exc:
            for item in exc.errors():
                field = str(item["loc"][0]) if item.get("loc") else "linha"
                label = FIELD_LABELS.get(field, field)
                if field == "email":
                    errors.append("E-mail inválido.")
                elif field == "legal_name":
                    errors.append("Razão social deve ter entre 2 e 200 caracteres.")
                elif field in {"trade_name", "tax_id", "phone"}:
                    errors.append(f"{label} excede o limite permitido ou é inválido.")
                else:
                    errors.append(f"{label}: dado inválido.")

        tax_key = canonical_tax_id(values.get("tax_id"))
        if tax_key:
            if tax_key in existing_tax_ids:
                errors.append("CNPJ/CPF já cadastrado neste escritório.")
            elif tax_key in seen_tax_ids:
                errors.append("CNPJ/CPF duplicado dentro do arquivo.")
            seen_tax_ids.add(tax_key)

        # Invalid cells are shortened in the stored preview; the source file itself is never persisted.
        for field in ("legal_name", "trade_name", "tax_id", "email", "phone"):
            value = values.get(field)
            if isinstance(value, str) and len(value) > 300:
                values[field] = value[:300]
        preview.append({"row_number": row_number, **values, "errors": sorted(set(errors))})
    return preview
