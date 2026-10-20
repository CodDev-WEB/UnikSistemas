import hashlib
import re
import secrets
from datetime import UTC, datetime, timedelta
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from io import BytesIO
from pathlib import PurePosixPath
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Header, HTTPException, Query, UploadFile, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, aliased

from app.config import Settings, get_settings
from app.db import get_db
from app.client_imports import canonical_tax_id, parse_client_import, validate_import_rows
from app.models import AntivirusHealth, AuditLog, Client, ClientDocument, ClientImportBatch, ClientTask, DocumentUploadLink, Notification, Tenant, User, utcnow
from app.schemas import (
    ClientCreate,
    ClientRead,
    ClientImportBatchRead,
    ClientImportHistoryRead,
    ClientTaskCreate,
    ClientTaskRead,
    ClientTaskUpdate,
    ClientDocumentRead,
    ClientDocumentReviewUpdate,
    DocumentUploadLinkRead,
    PublicDocumentUploadInfo,
    NotificationRead,
    LoginRequest,
    MeResponse,
    TenantBootstrapRequest,
    TenantBootstrapResponse,
    TokenResponse,
    UserCreate,
    UserUpdate,
    UserRead,
)
from app.security import create_access_token, get_current_user, hash_password, verify_password
from app.notifications import enqueue_due_reminders, notification_email_is_ready, notification_queue_counts, queue_notification
from app.document_storage import (
    ALLOWED_EXTENSIONS, LINK_LIFETIME_DAYS, MAX_DOCUMENT_BYTES, MAX_UPLOADS_PER_LINK,
    load_document, remove_stored_document, store_document, validate_uploaded_document,
)

router = APIRouter(prefix="/api/v1")


@router.post(
    "/setup/tenants",
    response_model=TenantBootstrapResponse,
    status_code=status.HTTP_201_CREATED,
    tags=["Setup"],
)
def bootstrap_tenant(
    payload: TenantBootstrapRequest,
    x_bootstrap_token: str | None = Header(default=None),
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Provision a tenant and its owner using an operator-held secret."""
    if not settings.bootstrap_token:
        raise HTTPException(status_code=503, detail="Provisionamento inicial está desabilitado.")
    if not x_bootstrap_token or not secrets.compare_digest(
        x_bootstrap_token, settings.bootstrap_token
    ):
        raise HTTPException(status_code=403, detail="Token de provisionamento inválido.")
    if db.scalar(select(Tenant).where(Tenant.slug == payload.tenant_slug)):
        raise HTTPException(status_code=409, detail="Já existe um escritório com esse identificador.")

    tenant = Tenant(name=payload.tenant_name, slug=payload.tenant_slug)
    try:
        db.add(tenant)
        db.flush()
        owner = User(
            tenant_id=tenant.id,
            email=str(payload.owner_email).lower(),
            full_name=payload.owner_full_name,
            password_hash=hash_password(payload.owner_password),
            role="owner",
        )
        db.add(owner)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Não foi possível criar o escritório.") from None

    db.refresh(tenant)
    db.refresh(owner)
    return TenantBootstrapResponse(tenant=tenant, owner=owner)


@router.post("/auth/token", response_model=TokenResponse, tags=["Authentication"])
def login(
    payload: LoginRequest,
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    tenant = db.scalar(select(Tenant).where(Tenant.slug == payload.tenant_slug))
    user = None
    if tenant and tenant.status == "active":
        user = db.scalar(
            select(User).where(
                User.tenant_id == tenant.id,
                User.email == str(payload.email).lower(),
                User.is_active.is_(True),
            )
        )
    if user is None or not verify_password(payload.password, user.password_hash):
        raise HTTPException(
            status_code=401,
            detail="Escritório, e-mail ou senha inválidos.",
            headers={"WWW-Authenticate": "Bearer"},
        )
    return TokenResponse(
        access_token=create_access_token(user, settings),
        expires_in=settings.access_token_expire_minutes * 60,
    )


@router.get("/me", response_model=MeResponse, tags=["Authentication"])
def read_me(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    tenant = db.scalar(select(Tenant).where(Tenant.id == current_user.tenant_id))
    if tenant is None:
        raise HTTPException(status_code=404, detail="Escritório não encontrado.")
    return MeResponse(user=current_user, tenant=tenant)


@router.get("/clients", response_model=list[ClientRead], tags=["Clients"])
def list_clients(
    offset: int = Query(default=0, ge=0),
    limit: int = Query(default=50, ge=1, le=100),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    return list(
        db.scalars(
            select(Client)
            .where(Client.tenant_id == current_user.tenant_id)
            .order_by(Client.created_at.desc(), Client.id)
            .offset(offset)
            .limit(limit)
        ).all()
    )


@router.post(
    "/clients",
    response_model=ClientRead,
    status_code=status.HTTP_201_CREATED,
    tags=["Clients"],
)
def create_client(
    payload: ClientCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if current_user.role not in {"owner", "admin", "accountant", "assistant"}:
        raise HTTPException(status_code=403, detail="Seu perfil não permite cadastrar clientes.")

    new_client = Client(
        tenant_id=current_user.tenant_id,
        legal_name=payload.legal_name,
        trade_name=payload.trade_name,
        tax_id=payload.tax_id,
        email=str(payload.email).lower() if payload.email else None,
        phone=payload.phone,
        status=payload.status,
    )
    try:
        db.add(new_client)
        db.flush()
        db.add(
            AuditLog(
                tenant_id=current_user.tenant_id,
                actor_user_id=current_user.id,
                action="client.created",
                entity_type="client",
                entity_id=new_client.id,
                details={"client_id": new_client.id},
            )
        )
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Já existe um cliente com esse identificador fiscal neste escritório.",
        ) from None
    db.refresh(new_client)
    return new_client




MANAGE_USER_ROLES = {"owner", "admin"}
VALID_STAFF_ROLES = {"admin", "accountant", "assistant", "viewer"}


def _require_user_manager(user: User) -> None:
    if user.role not in MANAGE_USER_ROLES:
        raise HTTPException(status_code=403, detail="Seu perfil não permite administrar usuários.")


def _record_user_audit(
    db: Session,
    actor: User,
    action: str,
    subject: User,
    changed_fields: list[str],
) -> None:
    db.add(
        AuditLog(
            tenant_id=actor.tenant_id,
            actor_user_id=actor.id,
            action=action,
            entity_type="user",
            entity_id=subject.id,
            details={"fields": sorted(changed_fields)},
        )
    )


@router.get("/users", response_model=list[UserRead], tags=["Users"])
def list_users(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    _require_user_manager(current_user)
    return list(
        db.scalars(
            select(User)
            .where(User.tenant_id == current_user.tenant_id)
            .order_by(User.role, User.full_name, User.id)
        ).all()
    )


@router.post(
    "/users",
    response_model=UserRead,
    status_code=status.HTTP_201_CREATED,
    tags=["Users"],
)
def create_user(
    payload: UserCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    _require_user_manager(current_user)
    if current_user.role == "admin" and payload.role == "admin":
        raise HTTPException(
            status_code=403,
            detail="Somente o responsável do escritório pode criar administradores.",
        )

    user = User(
        tenant_id=current_user.tenant_id,
        email=str(payload.email).lower(),
        full_name=payload.full_name,
        password_hash=hash_password(payload.password),
        role=payload.role,
        is_active=True,
    )
    try:
        db.add(user)
        db.flush()
        _record_user_audit(db, current_user, "user.created", user, ["email", "full_name", "role"])
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Já existe um usuário com esse e-mail neste escritório.",
        ) from None

    db.refresh(user)
    return user


@router.patch("/users/{user_id}", response_model=UserRead, tags=["Users"])
def update_user(
    user_id: str,
    payload: UserUpdate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    _require_user_manager(current_user)
    user = db.scalar(
        select(User).where(User.id == user_id, User.tenant_id == current_user.tenant_id)
    )
    if user is None:
        raise HTTPException(status_code=404, detail="Usuário não encontrado.")

    changes = payload.model_dump(exclude_unset=True)
    if not changes:
        raise HTTPException(status_code=422, detail="Informe ao menos um campo para atualizar.")
    if any(value is None for value in changes.values()):
        raise HTTPException(status_code=422, detail="Os campos enviados não podem ser nulos.")

    if user.id == current_user.id and ({"role", "is_active"} & changes.keys()):
        raise HTTPException(
            status_code=403,
            detail="Você não pode alterar o próprio perfil nem desativar o próprio acesso.",
        )
    if user.role == "owner" and user.id != current_user.id:
        raise HTTPException(
            status_code=403,
            detail="O usuário responsável não pode ser alterado por outro usuário.",
        )
    if current_user.role == "admin":
        if user.role == "admin" and user.id != current_user.id:
            raise HTTPException(
                status_code=403,
                detail="Administradores não podem alterar outros administradores.",
            )
        if changes.get("role") == "admin":
            raise HTTPException(
                status_code=403,
                detail="Somente o responsável do escritório pode atribuir o perfil administrador.",
            )

    old_active = user.is_active
    for field, value in changes.items():
        setattr(user, field, value)

    if "is_active" in changes and changes["is_active"] != old_active:
        action = "user.activated" if changes["is_active"] else "user.deactivated"
    else:
        action = "user.updated"

    try:
        _record_user_audit(db, current_user, action, user, list(changes.keys()))
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Não foi possível atualizar o usuário.") from None

    db.refresh(user)
    return user


@router.get("/clients/{client_id}", response_model=ClientRead, tags=["Clients"])
def get_client(
    client_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    client = db.scalar(
        select(Client).where(
            Client.id == client_id,
            Client.tenant_id == current_user.tenant_id,
        )
    )
    if client is None:
        raise HTTPException(status_code=404, detail="Cliente não encontrado.")
    return client



CLIENT_EDITOR_ROLES = {"owner", "admin", "accountant", "assistant"}
MAX_IMPORT_FILE_BYTES = 5 * 1024 * 1024


def _require_client_editor(user: User) -> None:
    if user.role not in CLIENT_EDITOR_ROLES:
        raise HTTPException(status_code=403, detail="Seu perfil não permite importar clientes.")


def _serialize_import_batch(batch: ClientImportBatch) -> dict:
    return {
        "id": batch.id,
        "filename": batch.filename,
        "status": batch.status,
        "total_rows": batch.total_rows,
        "imported_count": batch.imported_count,
        "rows": batch.preview_rows,
        "created_at": batch.created_at.isoformat() if batch.created_at else None,
        "confirmed_at": batch.confirmed_at.isoformat() if batch.confirmed_at else None,
    }


@router.post(
    "/imports/clients/preview",
    response_model=ClientImportBatchRead,
    status_code=status.HTTP_201_CREATED,
    tags=["Imports"],
)
async def preview_client_import(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Parse and validate a CSV/XLSX without creating any clients."""
    _require_client_editor(current_user)
    if not file.filename:
        raise HTTPException(status_code=422, detail="Selecione um arquivo para importar.")

    filename = PurePosixPath(file.filename.replace("\\", "/")).name[:255]
    try:
        content = await file.read(MAX_IMPORT_FILE_BYTES + 1)
        if len(content) > MAX_IMPORT_FILE_BYTES:
            raise HTTPException(
                status_code=413,
                detail="O arquivo ultrapassa o limite de 5 MiB.",
            )
        try:
            source_rows = parse_client_import(filename, content)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from None
    finally:
        await file.close()

    existing_ids = {
        key
        for value in db.scalars(
            select(Client.tax_id).where(
                Client.tenant_id == current_user.tenant_id,
                Client.tax_id.is_not(None),
            )
        ).all()
        if (key := canonical_tax_id(value))
    }
    preview_rows = validate_import_rows(source_rows, existing_ids)
    batch = ClientImportBatch(
        tenant_id=current_user.tenant_id,
        created_by_user_id=current_user.id,
        filename=filename,
        status="ready" if all(not row["errors"] for row in preview_rows) else "invalid",
        preview_rows=preview_rows,
        total_rows=len(preview_rows),
        imported_count=0,
    )
    db.add(batch)
    db.commit()
    db.refresh(batch)
    return _serialize_import_batch(batch)





@router.get(
    "/imports/clients",
    response_model=list[ClientImportHistoryRead],
    tags=["Imports"],
)
def list_client_import_history(
    offset: int = Query(default=0, ge=0),
    limit: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Return import batches for the currently authenticated office only."""
    _require_client_editor(current_user)
    records = db.execute(
        select(ClientImportBatch, User.full_name)
        .outerjoin(User, User.id == ClientImportBatch.created_by_user_id)
        .where(ClientImportBatch.tenant_id == current_user.tenant_id)
        .order_by(ClientImportBatch.created_at.desc(), ClientImportBatch.id.desc())
        .offset(offset)
        .limit(limit)
    ).all()
    return [
        {
            "id": batch.id,
            "filename": batch.filename,
            "status": batch.status,
            "total_rows": batch.total_rows,
            "imported_count": batch.imported_count,
            "created_by_name": creator_name,
            "created_at": batch.created_at.isoformat() if batch.created_at else None,
            "confirmed_at": batch.confirmed_at.isoformat() if batch.confirmed_at else None,
        }
        for batch, creator_name in records
    ]


@router.get("/imports/client-template.csv", tags=["Imports"])
def download_client_import_csv_template(
    current_user: User = Depends(get_current_user),
):
    """Download a minimal CSV template for client imports."""
    _require_client_editor(current_user)
    content = (
        "\ufeffRazão Social;Nome Fantasia;CNPJ/CPF;E-mail;Telefone;Status\r\n"
    )
    return Response(
        content=content.encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="unik-modelo-clientes.csv"'},
    )


@router.get("/imports/client-template.xlsx", tags=["Imports"])
def download_client_import_xlsx_template(
    current_user: User = Depends(get_current_user),
):
    """Download an Excel template with a clients tab and usage instructions."""
    _require_client_editor(current_user)
    workbook = Workbook()
    clients_sheet = workbook.active
    clients_sheet.title = "Clientes"
    headers = ["Razão Social", "Nome Fantasia", "CNPJ/CPF", "E-mail", "Telefone", "Status"]
    clients_sheet.append(headers)
    clients_sheet.freeze_panes = "A2"
    clients_sheet.auto_filter.ref = "A1:F1"
    widths = {"A": 34, "B": 28, "C": 20, "D": 32, "E": 20, "F": 16}
    for column, width in widths.items():
        clients_sheet.column_dimensions[column].width = width
    for cell in clients_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5C55B7")
    for row_number in range(2, 1002):
        clients_sheet.cell(row=row_number, column=3).number_format = "@"

    instructions = workbook.create_sheet("Instruções")
    instructions.append(["Como usar o modelo de clientes"])
    instructions.append(["1. Preencha uma linha para cada cliente na aba Clientes."])
    instructions.append(["2. Razão Social é obrigatória; mantenha a primeira linha de títulos."])
    instructions.append(["3. CNPJ/CPF deve ser formatado como texto para preservar zeros à esquerda."])
    instructions.append(["4. Status aceita Ativo ou Arquivado; se vazio, será considerado Ativo."])
    instructions.append(["5. Salve como .xlsx e envie pelo menu Importações para revisar a prévia."])
    instructions.column_dimensions["A"].width = 105
    instructions["A1"].font = Font(bold=True, size=14)
    workbook_buffer = BytesIO()
    workbook.save(workbook_buffer)
    workbook.close()
    return Response(
        content=workbook_buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="unik-modelo-clientes.xlsx"'},
    )


@router.get("/imports/clients/{batch_id}", response_model=ClientImportBatchRead, tags=["Imports"])
def read_client_import_preview(
    batch_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    _require_client_editor(current_user)
    batch = db.scalar(
        select(ClientImportBatch).where(
            ClientImportBatch.id == batch_id,
            ClientImportBatch.tenant_id == current_user.tenant_id,
        )
    )
    if batch is None:
        raise HTTPException(status_code=404, detail="Importação não encontrada.")
    return _serialize_import_batch(batch)


@router.delete(
    "/imports/clients/{batch_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    tags=["Imports"],
)
def discard_client_import(
    batch_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> None:
    """Mark a pending import preview as discarded and remove its personal data."""
    _require_client_editor(current_user)
    batch = db.scalar(
        select(ClientImportBatch).where(
            ClientImportBatch.id == batch_id,
            ClientImportBatch.tenant_id == current_user.tenant_id,
        )
    )
    if batch is None:
        raise HTTPException(status_code=404, detail="Importação não encontrada.")
    if batch.status == "confirmed":
        raise HTTPException(status_code=409, detail="A importação confirmada não pode ser descartada.")
    batch.status = "discarded"
    batch.preview_rows = []
    db.add(
        AuditLog(
            tenant_id=current_user.tenant_id,
            actor_user_id=current_user.id,
            action="clients.import.discarded",
            entity_type="client_import",
            entity_id=batch.id,
            details={"filename": batch.filename},
        )
    )
    db.commit()


@router.post(
    "/imports/clients/{batch_id}/confirm",
    response_model=ClientImportBatchRead,
    tags=["Imports"],
)
def confirm_client_import(
    batch_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Persist all rows atomically after a valid preview and explicit confirmation."""
    _require_client_editor(current_user)
    batch = db.scalar(
        select(ClientImportBatch)
        .where(
            ClientImportBatch.id == batch_id,
            ClientImportBatch.tenant_id == current_user.tenant_id,
        )
        .with_for_update()
    )
    if batch is None:
        raise HTTPException(status_code=404, detail="Importação não encontrada.")
    if batch.status == "confirmed":
        raise HTTPException(status_code=409, detail="Esta importação já foi confirmada.")
    if batch.status != "ready" or any(row.get("errors") for row in batch.preview_rows):
        raise HTTPException(
            status_code=409,
            detail="A prévia possui pendências. Corrija o arquivo e envie novamente antes de confirmar.",
        )

    existing_ids = {
        key
        for value in db.scalars(
            select(Client.tax_id).where(
                Client.tenant_id == current_user.tenant_id,
                Client.tax_id.is_not(None),
            )
        ).all()
        if (key := canonical_tax_id(value))
    }
    conflicts = [
        row for row in batch.preview_rows
        if (key := canonical_tax_id(row.get("tax_id"))) and key in existing_ids
    ]
    if conflicts:
        updated_rows = [dict(row) for row in batch.preview_rows]
        for row in updated_rows:
            key = canonical_tax_id(row.get("tax_id"))
            if key and key in existing_ids:
                row["errors"] = sorted(set([*row.get("errors", []), "CNPJ/CPF já cadastrado neste escritório."]))
        batch.preview_rows = updated_rows
        batch.status = "invalid"
        db.commit()
        raise HTTPException(
            status_code=409,
            detail="A base mudou após a prévia. Há CNPJ/CPF que já foi cadastrado; revise as linhas sinalizadas.",
        )

    try:
        clients = [
            Client(
                tenant_id=current_user.tenant_id,
                legal_name=row["legal_name"],
                trade_name=row.get("trade_name"),
                tax_id=row.get("tax_id"),
                email=row.get("email"),
                phone=row.get("phone"),
                status=row.get("status") or "active",
            )
            for row in batch.preview_rows
        ]
        db.add_all(clients)
        db.flush()
        batch.status = "confirmed"
        batch.imported_count = len(clients)
        batch.confirmed_at = utcnow()
        # Client data now lives in the clients table; remove the copy stored for preview.
        batch.preview_rows = []
        db.add(
            AuditLog(
                tenant_id=current_user.tenant_id,
                actor_user_id=current_user.id,
                action="clients.imported",
                entity_type="client_import",
                entity_id=batch.id,
                details={"imported_count": len(clients), "filename": batch.filename},
            )
        )
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Um ou mais identificadores fiscais já existem. Nenhum cliente foi importado; envie a base atualizada.",
        ) from None
    db.refresh(batch)
    return _serialize_import_batch(batch)



TASK_EDITOR_ROLES = {"owner", "admin", "accountant", "assistant"}


def _require_task_editor(user: User) -> None:
    if user.role not in TASK_EDITOR_ROLES:
        raise HTTPException(status_code=403, detail="Seu perfil permite apenas consultar as pendências.")


def _task_snapshot_query():
    assigned_user = aliased(User)
    creator_user = aliased(User)
    return (
        select(
            ClientTask,
            Client.legal_name.label("client_name"),
            assigned_user.full_name.label("assigned_to_name"),
            creator_user.full_name.label("created_by_name"),
        )
        .join(Client, Client.id == ClientTask.client_id)
        .outerjoin(assigned_user, assigned_user.id == ClientTask.assigned_to_user_id)
        .outerjoin(creator_user, creator_user.id == ClientTask.created_by_user_id)
    )


def _serialize_task_record(record) -> dict:
    task, client_name, assigned_to_name, created_by_name = record
    return {
        "id": task.id,
        "client_id": task.client_id,
        "client_name": client_name,
        "title": task.title,
        "description": task.description,
        "kind": task.kind,
        "priority": task.priority,
        "status": task.status,
        "due_date": task.due_date,
        "assigned_to_user_id": task.assigned_to_user_id,
        "assigned_to_name": assigned_to_name,
        "created_by_user_id": task.created_by_user_id,
        "created_by_name": created_by_name,
        "completed_at": task.completed_at.isoformat() if task.completed_at else None,
        "created_at": task.created_at.isoformat() if task.created_at else None,
        "updated_at": task.updated_at.isoformat() if task.updated_at else None,
    }


def _read_task(db: Session, task_id: str, tenant_id: str):
    record = db.execute(
        _task_snapshot_query().where(
            ClientTask.id == task_id,
            ClientTask.tenant_id == tenant_id,
        )
    ).one_or_none()
    if record is None:
        raise HTTPException(status_code=404, detail="Pendência não encontrada.")
    return record


def _validate_task_assignment(db: Session, user_id: str | None, tenant_id: str) -> User | None:
    if user_id is None:
        return None
    assigned = db.scalar(
        select(User).where(
            User.id == user_id,
            User.tenant_id == tenant_id,
            User.is_active.is_(True),
        )
    )
    if assigned is None:
        raise HTTPException(
            status_code=422,
            detail="O responsável precisa ser um usuário ativo deste escritório.",
        )
    return assigned


@router.get("/tasks", response_model=list[ClientTaskRead], tags=["Tasks"])
def list_client_tasks(
    status_filter: str | None = Query(
        default=None,
        alias="status",
        pattern="^(open|received|completed|cancelled)$",
    ),
    client_id: str | None = Query(default=None, min_length=1, max_length=36),
    limit: int = Query(default=100, ge=1, le=100),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """List tasks and document requests scoped to the authenticated office."""
    statement = _task_snapshot_query().where(ClientTask.tenant_id == current_user.tenant_id)
    if status_filter is not None:
        statement = statement.where(ClientTask.status == status_filter)
    if client_id is not None:
        statement = statement.where(
            ClientTask.client_id == client_id,
            Client.tenant_id == current_user.tenant_id,
        )
    records = db.execute(
        statement.order_by(
            ClientTask.due_date.is_(None),
            ClientTask.due_date.asc(),
            ClientTask.created_at.desc(),
        ).limit(limit)
    ).all()
    return [_serialize_task_record(record) for record in records]


@router.post(
    "/tasks",
    response_model=ClientTaskRead,
    status_code=status.HTTP_201_CREATED,
    tags=["Tasks"],
)
def create_client_task(
    payload: ClientTaskCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    _require_task_editor(current_user)
    target_client = db.scalar(
        select(Client).where(
            Client.id == payload.client_id,
            Client.tenant_id == current_user.tenant_id,
        )
    )
    if target_client is None:
        raise HTTPException(status_code=404, detail="Cliente não encontrado neste escritório.")

    assignee_id = payload.assigned_to_user_id or current_user.id
    assigned_user = _validate_task_assignment(db, assignee_id, current_user.tenant_id)
    task = ClientTask(
        tenant_id=current_user.tenant_id,
        client_id=target_client.id,
        created_by_user_id=current_user.id,
        assigned_to_user_id=assignee_id,
        title=payload.title,
        description=payload.description,
        kind=payload.kind,
        priority=payload.priority,
        due_date=payload.due_date,
        status="open",
    )
    try:
        db.add(task)
        db.flush()
        if assigned_user is not None and assigned_user.id != current_user.id:
            queue_notification(
                db,
                tenant_id=current_user.tenant_id,
                recipient_user_id=assigned_user.id,
                event_key=f"task_assigned:{task.id}",
                event_type="task_assigned",
                title="Nova pendência atribuída",
                message=f"A pendência “{task.title}” foi atribuída a você para o cliente {target_client.legal_name}.",
                entity_type="client_task",
                entity_id=task.id,
                action_url="/tasks",
            )
        db.add(
            AuditLog(
                tenant_id=current_user.tenant_id,
                actor_user_id=current_user.id,
                action="task.created",
                entity_type="client_task",
                entity_id=task.id,
                details={
                    "client_id": task.client_id,
                    "kind": task.kind,
                    "due_date": task.due_date.isoformat() if task.due_date else None,
                },
            )
        )
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Não foi possível registrar a pendência.") from None
    return _serialize_task_record(_read_task(db, task.id, current_user.tenant_id))


@router.patch("/tasks/{task_id}", response_model=ClientTaskRead, tags=["Tasks"])
def update_client_task(
    task_id: str,
    payload: ClientTaskUpdate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    _require_task_editor(current_user)
    task = db.scalar(
        select(ClientTask).where(
            ClientTask.id == task_id,
            ClientTask.tenant_id == current_user.tenant_id,
        )
    )
    if task is None:
        raise HTTPException(status_code=404, detail="Pendência não encontrada.")
    changes = payload.model_dump(exclude_unset=True)
    if not changes:
        raise HTTPException(status_code=422, detail="Informe ao menos um campo para atualizar.")
    nullable_fields = {"description", "due_date", "assigned_to_user_id"}
    if any(value is None and field not in nullable_fields for field, value in changes.items()):
        raise HTTPException(status_code=422, detail="Campos obrigatórios não podem ser nulos.")

    new_assignee = None
    if "assigned_to_user_id" in changes and changes["assigned_to_user_id"] is not None:
        new_assignee = _validate_task_assignment(db, changes["assigned_to_user_id"], current_user.tenant_id)
    previous_assignee_id = task.assigned_to_user_id

    previous_status = task.status
    previous_due_date = task.due_date
    for field, value in changes.items():
        setattr(task, field, value)
    task.updated_at = utcnow()
    if "status" in changes and changes["status"] != previous_status:
        task.completed_at = utcnow() if changes["status"] == "completed" else None

    try:
        if (
            "assigned_to_user_id" in changes
            and task.assigned_to_user_id != previous_assignee_id
            and new_assignee is not None
            and new_assignee.id != current_user.id
        ):
            queue_notification(
                db,
                tenant_id=current_user.tenant_id,
                recipient_user_id=new_assignee.id,
                event_key=f"task_assigned:{task.id}:{new_assignee.id}:{task.updated_at.isoformat()}",
                event_type="task_assigned",
                title="Nova pendência atribuída",
                message=f"A pendência “{task.title}” foi atribuída a você.",
                entity_type="client_task",
                entity_id=task.id,
                action_url="/tasks",
            )
        db.add(
            AuditLog(
                tenant_id=current_user.tenant_id,
                actor_user_id=current_user.id,
                action="task.updated",
                entity_type="client_task",
                entity_id=task.id,
                details={
                    "fields": sorted(changes),
                    "previous_status": previous_status,
                    "status": task.status,
                    "previous_due_date": previous_due_date.isoformat() if previous_due_date else None,
                    "due_date": task.due_date.isoformat() if task.due_date else None,
                },
            )
        )
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Não foi possível atualizar a pendência.") from None
    return _serialize_task_record(_read_task(db, task.id, current_user.tenant_id))


def _token_digest(token: str) -> str:
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def _datetime_expired(value) -> bool:
    if value.tzinfo is None:
        value = value.replace(tzinfo=utcnow().tzinfo)
    return value <= utcnow()


def _resolve_upload_link(db: Session, token: str, lock: bool = False):
    if not token or len(token) > 256:
        return None
    query = select(DocumentUploadLink).where(
        DocumentUploadLink.token_hash == _token_digest(token)
    )
    if lock:
        query = query.with_for_update().execution_options(populate_existing=True)
    link = db.scalar(query)
    if (
        link is None
        or link.revoked_at is not None
        or _datetime_expired(link.expires_at)
        or link.upload_count >= link.max_uploads
    ):
        return None
    task = db.scalar(
        select(ClientTask).where(
            ClientTask.id == link.task_id,
            ClientTask.tenant_id == link.tenant_id,
        )
    )
    if task is None or task.kind != "document" or task.status in {"completed", "cancelled"}:
        return None
    return link, task


def _serialize_document_record(record) -> dict:
    document, task_title, client_name, reviewed_by_name = record
    return {
        "id": document.id,
        "task_id": document.task_id,
        "task_title": task_title,
        "client_name": client_name,
        "original_filename": document.original_filename,
        "content_type": document.content_type,
        "file_size": document.file_size,
        "sha256": document.sha256,
        "scan_status": document.scan_status,
        "scan_attempts": document.scan_attempts,
        "scan_completed_at": document.scan_completed_at.isoformat() if document.scan_completed_at else None,
        "review_status": document.review_status,
        "review_note": document.review_note,
        "reviewed_by_name": reviewed_by_name,
        "reviewed_at": document.reviewed_at.isoformat() if document.reviewed_at else None,
        "created_at": document.created_at.isoformat() if document.created_at else None,
    }


def _document_snapshot_query():
    reviewer = aliased(User)
    return (
        select(
            ClientDocument,
            ClientTask.title.label("task_title"),
            Client.legal_name.label("client_name"),
            reviewer.full_name.label("reviewed_by_name"),
        )
        .join(
            ClientTask,
            (ClientTask.id == ClientDocument.task_id)
            & (ClientTask.tenant_id == ClientDocument.tenant_id),
        )
        .join(Client, (Client.id == ClientTask.client_id) & (Client.tenant_id == ClientTask.tenant_id))
        .outerjoin(reviewer, reviewer.id == ClientDocument.reviewed_by_user_id)
    )


@router.post(
    "/tasks/{task_id}/upload-link",
    response_model=DocumentUploadLinkRead,
    status_code=status.HTTP_201_CREATED,
    tags=["Documents"],
)
def create_document_upload_link(
    task_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    _require_task_editor(current_user)
    task = db.scalar(
        select(ClientTask).where(
            ClientTask.id == task_id,
            ClientTask.tenant_id == current_user.tenant_id,
        )
    )
    if task is None:
        raise HTTPException(status_code=404, detail="Pendência não encontrada.")
    if task.kind != "document" or task.status in {"completed", "cancelled"}:
        raise HTTPException(
            status_code=409,
            detail="O link só pode ser criado para uma solicitação de documento em aberto.",
        )

    now = utcnow()
    previous_links = db.scalars(
        select(DocumentUploadLink).where(
            DocumentUploadLink.tenant_id == current_user.tenant_id,
            DocumentUploadLink.task_id == task.id,
            DocumentUploadLink.revoked_at.is_(None),
        )
    ).all()
    for previous_link in previous_links:
        previous_link.revoked_at = now
    raw_token = secrets.token_urlsafe(32)
    expiry = now + timedelta(days=LINK_LIFETIME_DAYS)
    link = DocumentUploadLink(
        tenant_id=current_user.tenant_id,
        task_id=task.id,
        created_by_user_id=current_user.id,
        token_hash=_token_digest(raw_token),
        expires_at=expiry,
        max_uploads=MAX_UPLOADS_PER_LINK,
        upload_count=0,
    )
    db.add(link)
    db.add(
        AuditLog(
            tenant_id=current_user.tenant_id,
            actor_user_id=current_user.id,
            action="document_upload_link.created",
            entity_type="client_task",
            entity_id=task.id,
            details={"expires_at": expiry.isoformat(), "max_uploads": MAX_UPLOADS_PER_LINK},
        )
    )
    db.commit()
    return DocumentUploadLinkRead(
        id=link.id,
        public_path=f"/document-upload/{raw_token}",
        expires_at=expiry.isoformat(),
        max_uploads=MAX_UPLOADS_PER_LINK,
    )


@router.post(
    "/tasks/{task_id}/upload-links/{link_id}/revoke",
    status_code=status.HTTP_204_NO_CONTENT,
    tags=["Documents"],
)
def revoke_document_upload_link(
    task_id: str,
    link_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    _require_task_editor(current_user)
    link = db.scalar(
        select(DocumentUploadLink).where(
            DocumentUploadLink.id == link_id,
            DocumentUploadLink.task_id == task_id,
            DocumentUploadLink.tenant_id == current_user.tenant_id,
        )
    )
    if link is None:
        raise HTTPException(status_code=404, detail="Link de envio não encontrado.")
    if link.revoked_at is None:
        link.revoked_at = utcnow()
        db.add(
            AuditLog(
                tenant_id=current_user.tenant_id,
                actor_user_id=current_user.id,
                action="document_upload_link.revoked",
                entity_type="client_task",
                entity_id=task_id,
                details={"upload_link_id": link.id},
            )
        )
        db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get(
    "/public/document-upload/{token}",
    response_model=PublicDocumentUploadInfo,
    tags=["Public document intake"],
)
def read_public_document_upload_link(token: str, db: Session = Depends(get_db)):
    resolved = _resolve_upload_link(db, token)
    if resolved is None:
        raise HTTPException(status_code=404, detail="Este link não está disponível.")
    link, task = resolved
    return PublicDocumentUploadInfo(
        title=task.title,
        description=task.description,
        expires_at=link.expires_at.isoformat(),
        remaining_uploads=max(0, link.max_uploads - link.upload_count),
        max_file_size_bytes=MAX_DOCUMENT_BYTES,
        allowed_extensions=list(ALLOWED_EXTENSIONS),
    )


@router.post(
    "/public/document-upload/{token}/files",
    status_code=status.HTTP_201_CREATED,
    tags=["Public document intake"],
)
async def upload_public_document(
    token: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    if _resolve_upload_link(db, token) is None:
        raise HTTPException(status_code=404, detail="Este link não está disponível.")
    content = await file.read(MAX_DOCUMENT_BYTES + 1)
    if len(content) > MAX_DOCUMENT_BYTES:
        raise HTTPException(status_code=413, detail="O arquivo ultrapassa o limite de 10 MiB.")
    try:
        clean_name, content_type = validate_uploaded_document(file.filename, content)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from None

    resolved = _resolve_upload_link(db, token, lock=True)
    if resolved is None:
        raise HTTPException(status_code=404, detail="Este link não está mais disponível.")
    link, task = resolved
    extension = PurePosixPath(clean_name).suffix.lower()
    storage_key = store_document(settings.document_storage_path, extension, content)
    document = ClientDocument(
        tenant_id=link.tenant_id,
        task_id=task.id,
        upload_link_id=link.id,
        original_filename=clean_name,
        storage_key=storage_key,
        content_type=content_type,
        file_size=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        review_status="pending",
        scan_status="pending",
    )
    link.upload_count += 1
    if task.status == "open":
        task.status = "received"
    task.updated_at = utcnow()
    db.add(document)
    try:
        # Assign the ORM-generated document ID before writing the audit event.
        db.flush()
        recipient_ids = {user_id for user_id in (task.assigned_to_user_id, task.created_by_user_id) if user_id}
        for recipient_id in recipient_ids:
            recipient = db.scalar(
                select(User).where(
                    User.id == recipient_id,
                    User.tenant_id == link.tenant_id,
                    User.is_active.is_(True),
                )
            )
            if recipient is not None:
                queue_notification(
                    db,
                    tenant_id=link.tenant_id,
                    recipient_user_id=recipient.id,
                    event_key=f"document_received:{document.id}",
                    event_type="document_received",
                    title="Documento recebido",
                    message=f'O arquivo “{clean_name}” foi recebido para a pendência “{task.title}” e aguarda análise antimalware antes de ser liberado.',
                    entity_type="client_document",
                    entity_id=document.id,
                    action_url="/tasks",
                )
        db.add(
            AuditLog(
                tenant_id=link.tenant_id,
                actor_user_id=None,
                action="document.uploaded",
                entity_type="client_document",
                entity_id=document.id,
                details={
                    "task_id": task.id,
                    "file_size": len(content),
                    "content_type": content_type,
                },
            )
        )
        db.commit()
        db.refresh(document)
    except Exception:
        db.rollback()
        remove_stored_document(settings.document_storage_path, storage_key)
        raise
    return {
        "id": document.id,
        "filename": document.original_filename,
        "file_size": document.file_size,
        "review_status": document.review_status,
        "scan_status": document.scan_status,
    }


@router.get(
    "/documents",
    response_model=list[ClientDocumentRead],
    tags=["Documents"],
)
def list_client_documents(
    task_id: str | None = Query(default=None, min_length=1, max_length=36),
    limit: int = Query(default=100, ge=1, le=100),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    statement = _document_snapshot_query().where(
        ClientDocument.tenant_id == current_user.tenant_id
    )
    if task_id is not None:
        task_exists = db.scalar(
            select(ClientTask.id).where(
                ClientTask.id == task_id,
                ClientTask.tenant_id == current_user.tenant_id,
            )
        )
        if task_exists is None:
            raise HTTPException(status_code=404, detail="Pendência não encontrada.")
        statement = statement.where(ClientDocument.task_id == task_id)
    records = db.execute(
        statement.order_by(ClientDocument.created_at.desc()).limit(limit)
    ).all()
    return [_serialize_document_record(record) for record in records]


@router.get("/documents/{document_id}/download", tags=["Documents"])
def download_client_document(
    document_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    document = db.scalar(
        select(ClientDocument).where(
            ClientDocument.id == document_id,
            ClientDocument.tenant_id == current_user.tenant_id,
        )
    )
    if document is None:
        raise HTTPException(status_code=404, detail="Documento não encontrado.")
    if document.scan_status != "clean":
        raise HTTPException(
            status_code=423,
            detail="O documento permanece em quarentena até a verificação antimalware ser aprovada.",
        )
    try:
        content = load_document(settings.document_storage_path, document.storage_key, area="clean")
    except (FileNotFoundError, OSError):
        raise HTTPException(status_code=404, detail="O arquivo não está disponível no armazenamento.") from None

    db.add(
        AuditLog(
            tenant_id=current_user.tenant_id,
            actor_user_id=current_user.id,
            action="document.downloaded",
            entity_type="client_document",
            entity_id=document.id,
            details={"file_size": document.file_size},
        )
    )
    db.commit()
    quoted_name = quote(document.original_filename, safe="")
    ascii_name = re.sub(r"[^A-Za-z0-9._-]", "_", document.original_filename)
    return Response(
        content=content,
        media_type=document.content_type,
        headers={
            "Content-Disposition": f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{quoted_name}',
            "X-Content-Type-Options": "nosniff",
            "Cache-Control": "private, no-store",
        },
    )


@router.patch(
    "/documents/{document_id}",
    response_model=ClientDocumentRead,
    tags=["Documents"],
)
def review_client_document(
    document_id: str,
    payload: ClientDocumentReviewUpdate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    _require_task_editor(current_user)
    document = db.scalar(
        select(ClientDocument).where(
            ClientDocument.id == document_id,
            ClientDocument.tenant_id == current_user.tenant_id,
        )
    )
    if document is None:
        raise HTTPException(status_code=404, detail="Documento não encontrado.")
    if document.scan_status != "clean":
        raise HTTPException(
            status_code=409,
            detail="A conferência manual só fica disponível após a análise antimalware aprovada.",
        )
    document.review_status = payload.review_status
    document.review_note = payload.review_note
    document.reviewed_by_user_id = current_user.id
    document.reviewed_at = utcnow()
    db.add(
        AuditLog(
            tenant_id=current_user.tenant_id,
            actor_user_id=current_user.id,
            action="document.reviewed",
            entity_type="client_document",
            entity_id=document.id,
            details={"review_status": document.review_status},
        )
    )
    db.commit()
    record = db.execute(
        _document_snapshot_query().where(
            ClientDocument.id == document.id,
            ClientDocument.tenant_id == current_user.tenant_id,
        )
    ).one()
    return _serialize_document_record(record)


@router.post(
    "/documents/{document_id}/retry-scan",
    response_model=ClientDocumentRead,
    tags=["Documents"],
)
def retry_document_scan(
    document_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if current_user.role not in {"owner", "admin"}:
        raise HTTPException(status_code=403, detail="Somente o responsável ou administrador pode reenfileirar a análise.")
    document = db.scalar(
        select(ClientDocument).where(
            ClientDocument.id == document_id,
            ClientDocument.tenant_id == current_user.tenant_id,
        )
    )
    if document is None:
        raise HTTPException(status_code=404, detail="Documento não encontrado.")
    if document.scan_status != "error":
        raise HTTPException(
            status_code=409,
            detail="Somente documentos cuja análise falhou podem ser reenfileirados.",
        )
    document.scan_status = "pending"
    document.scan_attempts = 0
    document.scan_last_attempt_at = None
    document.scan_completed_at = None
    document.scan_error = None
    db.add(
        AuditLog(
            tenant_id=current_user.tenant_id,
            actor_user_id=current_user.id,
            action="document.scan_retry_requested",
            entity_type="client_document",
            entity_id=document.id,
            details={},
        )
    )
    db.commit()
    record = db.execute(
        _document_snapshot_query().where(
            ClientDocument.id == document.id,
            ClientDocument.tenant_id == current_user.tenant_id,
        )
    ).one()
    return _serialize_document_record(record)


def _serialize_notification(notification: Notification) -> dict:
    return {
        "id": notification.id,
        "event_type": notification.event_type,
        "title": notification.title,
        "message": notification.message,
        "entity_type": notification.entity_type,
        "entity_id": notification.entity_id,
        "action_url": notification.action_url,
        "is_read": notification.is_read,
        "email_status": notification.email_status,
        "email_attempts": notification.email_attempts,
        "email_sent_at": notification.email_sent_at.isoformat() if notification.email_sent_at else None,
        "created_at": notification.created_at.isoformat() if notification.created_at else None,
    }


def _require_notification_manager(user: User) -> None:
    if user.role not in {"owner", "admin"}:
        raise HTTPException(
            status_code=403,
            detail="Apenas o responsável ou administrador pode consultar e gerenciar a fila de e-mails.",
        )


@router.get("/security/antivirus-status", tags=["Security"])
def get_antivirus_status(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Expose scanner health and quarantine counts only to office owners/admins."""
    _require_notification_manager(current_user)
    health = db.get(AntivirusHealth, "clamav")
    status_value = health.status if health is not None else "unknown"
    database_date = health.signature_database_date if health is not None else None
    signature_age_hours = None
    if database_date is not None:
        current = datetime.now(UTC)
        built = (
            database_date.replace(tzinfo=UTC)
            if database_date.tzinfo is None
            else database_date.astimezone(UTC)
        )
        signature_age_hours = round((current - built).total_seconds() / 3600, 1)
    pending = int(
        db.scalar(
            select(func.count(ClientDocument.id)).where(
                ClientDocument.tenant_id == current_user.tenant_id,
                ClientDocument.scan_status.in_(("pending", "scanning"))
            )
        )
        or 0
    )
    failed = int(
        db.scalar(
            select(func.count(ClientDocument.id)).where(
                ClientDocument.tenant_id == current_user.tenant_id,
                ClientDocument.scan_status == "error",
            )
        )
        or 0
    )
    infected_quarantine = int(
        db.scalar(
            select(func.count(ClientDocument.id)).where(
                ClientDocument.tenant_id == current_user.tenant_id,
                ClientDocument.scan_status == "infected",
                ClientDocument.quarantine_purged_at.is_(None),
            )
        )
        or 0
    )
    return {
        "status": status_value,
        "checked_at": health.checked_at.isoformat() if health and health.checked_at else None,
        "last_success_at": (
            health.last_success_at.isoformat() if health and health.last_success_at else None
        ),
        "signature_database_version": health.signature_database_version if health else None,
        "signature_database_date": database_date.isoformat() if database_date else None,
        "signature_age_hours": signature_age_hours,
        "signature_max_age_hours": settings.antivirus_signature_max_age_hours,
        "last_error": health.last_error if health else None,
        "pending_documents": pending,
        "failed_documents": failed,
        "infected_quarantine_documents": infected_quarantine,
        "infected_retention_days": settings.antivirus_infected_retention_days,
        "last_quarantine_cleanup_at": (
            health.last_cleanup_at.isoformat() if health and health.last_cleanup_at else None
        ),
    }


@router.get(
    "/notifications",
    response_model=list[NotificationRead],
    tags=["Notifications"],
)
def list_notifications(
    unread_only: bool = False,
    limit: int = Query(default=50, ge=1, le=100),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    statement = select(Notification).where(
        Notification.tenant_id == current_user.tenant_id,
        Notification.recipient_user_id == current_user.id,
    )
    if unread_only:
        statement = statement.where(Notification.is_read.is_(False))
    rows = db.scalars(
        statement.order_by(Notification.created_at.desc(), Notification.id.desc()).limit(limit)
    ).all()
    return [_serialize_notification(row) for row in rows]


@router.patch(
    "/notifications/{notification_id}/read",
    response_model=NotificationRead,
    tags=["Notifications"],
)
def mark_notification_read(
    notification_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    notification = db.scalar(
        select(Notification).where(
            Notification.id == notification_id,
            Notification.tenant_id == current_user.tenant_id,
            Notification.recipient_user_id == current_user.id,
        )
    )
    if notification is None:
        raise HTTPException(status_code=404, detail="Notificação não encontrada.")
    if not notification.is_read:
        notification.is_read = True
        notification.read_at = utcnow()
        db.commit()
        db.refresh(notification)
    return _serialize_notification(notification)


@router.post("/notifications/read-all", tags=["Notifications"])
def mark_all_notifications_read(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    notifications = db.scalars(
        select(Notification).where(
            Notification.tenant_id == current_user.tenant_id,
            Notification.recipient_user_id == current_user.id,
            Notification.is_read.is_(False),
        )
    ).all()
    now = utcnow()
    for notification in notifications:
        notification.is_read = True
        notification.read_at = now
    db.commit()
    return {"marked_read": len(notifications)}


@router.get("/notifications/email-status", tags=["Notifications"])
def get_notification_email_status(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    _require_notification_manager(current_user)
    counts = notification_queue_counts(db, current_user.tenant_id)
    return {
        "enabled": settings.email_notifications_enabled,
        "configured": notification_email_is_ready(settings),
        "pending": counts["pending"],
        "failed": counts["failed"],
        "expired": counts["expired"],
        "max_queue_age_hours": settings.email_queue_max_age_hours,
        "note": (
            f"O worker envia automaticamente apenas mensagens na fila há até {settings.email_queue_max_age_hours} horas. "
            "Mensagens mais antigas expiram e exigem reenvio explícito por um responsável."
            if notification_email_is_ready(settings)
            else (
                "SMTP está desativado; notificações permanecem na caixa interna. "
                f"Ao habilitar, só serão enviadas automaticamente mensagens enfileiradas há até "
                f"{settings.email_queue_max_age_hours} horas; mensagens antigas exigem reenvio explícito."
            )
        ),
    }


@router.post("/notifications/process-due-reminders", tags=["Notifications"])
def process_due_reminders_now(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    _require_notification_manager(current_user)
    try:
        today = datetime.now(ZoneInfo(settings.notification_timezone)).date()
    except ZoneInfoNotFoundError:
        raise HTTPException(status_code=500, detail="Fuso horário de notificações inválido.") from None
    result = enqueue_due_reminders(db, today, days_ahead=1, tenant_id=current_user.tenant_id)
    db.commit()
    counts = notification_queue_counts(db, current_user.tenant_id)
    return {
        "created": result["created"],
        "due_items_checked": result["checked"],
        "email_enabled": notification_email_is_ready(settings),
        "queued_email_count": counts["pending"],
        "failed_email_count": counts["failed"],
    }


@router.post("/notifications/{notification_id}/retry-email", tags=["Notifications"])
def retry_notification_email(
    notification_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    _require_notification_manager(current_user)
    if not notification_email_is_ready(settings):
        raise HTTPException(status_code=409, detail="Configure SMTP e habilite o envio de e-mails antes de tentar novamente.")
    notification = db.scalar(
        select(Notification).where(
            Notification.id == notification_id,
            Notification.tenant_id == current_user.tenant_id,
        )
    )
    if notification is None:
        raise HTTPException(status_code=404, detail="Notificação não encontrada neste escritório.")
    if notification.email_status == "sent":
        raise HTTPException(status_code=409, detail="Este e-mail já foi enviado.")
    notification.email_status = "pending"
    notification.email_queued_at = utcnow()
    notification.email_last_error = None
    db.add(
        AuditLog(
            tenant_id=current_user.tenant_id,
            actor_user_id=current_user.id,
            action="notification.email_retry_requested",
            entity_type="notification",
            entity_id=notification.id,
            details={"reason": "explicit_manual_retry"},
        )
    )
    db.commit()
    return {"queued": True, "id": notification.id}



@router.post("/notifications/retry-failed-emails", tags=["Notifications"])
def retry_failed_notification_emails(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    _require_notification_manager(current_user)
    if not notification_email_is_ready(settings):
        raise HTTPException(status_code=409, detail="Configure SMTP e habilite o envio de e-mails antes de tentar novamente.")
    failed = db.scalars(
        select(Notification).where(
            Notification.tenant_id == current_user.tenant_id,
            Notification.email_status == "failed",
        )
    ).all()
    for notification in failed:
        notification.email_status = "pending"
        notification.email_queued_at = utcnow()
        notification.email_last_error = None
        db.add(
            AuditLog(
                tenant_id=current_user.tenant_id,
                actor_user_id=current_user.id,
                action="notification.email_retry_requested",
                entity_type="notification",
                entity_id=notification.id,
                details={"reason": "explicit_manual_retry_all"},
            )
        )
    db.commit()
    return {"queued": len(failed)}
