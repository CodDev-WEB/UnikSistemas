from fastapi.testclient import TestClient


def provision(client: TestClient, slug: str, email: str) -> dict:
    response = client.post(
        "/api/v1/setup/tenants",
        headers={"X-Bootstrap-Token": "test-bootstrap-token-with-more-than-thirty-two"},
        json={
            "tenant_name": f"Escritório {slug}",
            "tenant_slug": slug,
            "owner_email": email,
            "owner_full_name": "Responsável Contábil",
            "owner_password": "Senha-Forte-Teste-2026",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def login(client: TestClient, slug: str, email: str) -> str:
    response = client.post(
        "/api/v1/auth/token",
        json={
            "tenant_slug": slug,
            "email": email,
            "password": "Senha-Forte-Teste-2026",
        },
    )
    assert response.status_code == 200, response.text
    return response.json()["access_token"]


def test_health_endpoints(client: TestClient):
    assert client.get("/health/live").json()["status"] == "ok"
    ready = client.get("/health/ready")
    assert ready.status_code == 200
    assert ready.json()["database"] == "ok"


def test_bootstrap_requires_operator_token(client: TestClient):
    response = client.post(
        "/api/v1/setup/tenants",
        json={
            "tenant_name": "Escritório Teste",
            "tenant_slug": "teste",
            "owner_email": "owner@example.com",
            "owner_full_name": "Dono do Escritório",
            "owner_password": "Senha-Forte-Teste-2026",
        },
    )
    assert response.status_code == 403


def test_can_provision_login_and_create_client(client: TestClient):
    provision(client, "escritorio-teste", "owner@example.com")
    token = login(client, "escritorio-teste", "owner@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    me = client.get("/api/v1/me", headers=headers)
    assert me.status_code == 200
    assert me.json()["tenant"]["slug"] == "escritorio-teste"
    assert "password_hash" not in me.json()["user"]

    created = client.post(
        "/api/v1/clients",
        headers=headers,
        json={
            "legal_name": "Empresa Exemplo LTDA",
            "trade_name": "Empresa Exemplo",
            "tax_id": "12.345.678/0001-99",
            "email": "financeiro@example.com",
        },
    )
    assert created.status_code == 201, created.text
    client_id = created.json()["id"]

    listed = client.get("/api/v1/clients", headers=headers)
    assert listed.status_code == 200
    assert len(listed.json()) == 1
    assert listed.json()[0]["id"] == client_id


def test_tenant_cannot_read_another_tenants_client(client: TestClient):
    provision(client, "escritorio-a", "owner-a@example.com")
    token_a = login(client, "escritorio-a", "owner-a@example.com")
    created = client.post(
        "/api/v1/clients",
        headers={"Authorization": f"Bearer {token_a}"},
        json={"legal_name": "Cliente Privado do A"},
    )
    assert created.status_code == 201
    client_id = created.json()["id"]

    provision(client, "escritorio-b", "owner-b@example.com")
    token_b = login(client, "escritorio-b", "owner-b@example.com")
    response = client.get(
        f"/api/v1/clients/{client_id}",
        headers={"Authorization": f"Bearer {token_b}"},
    )
    assert response.status_code == 404
    assert response.json()["detail"] == "Cliente não encontrado."


def test_invalid_login_is_rejected(client: TestClient):
    provision(client, "escritorio-teste", "owner@example.com")
    response = client.post(
        "/api/v1/auth/token",
        json={
            "tenant_slug": "escritorio-teste",
            "email": "owner@example.com",
            "password": "senha-errada",
        },
    )
    assert response.status_code == 401


def test_endpoints_require_authentication(client: TestClient):
    assert client.get("/api/v1/clients").status_code == 401



def create_user(
    client: TestClient,
    token: str,
    email: str,
    role: str = "accountant",
    full_name: str = "Pessoa da Equipe",
) -> dict:
    response = client.post(
        "/api/v1/users",
        headers={"Authorization": f"Bearer {token}"},
        json={
            "email": email,
            "full_name": full_name,
            "password": "Senha-Forte-Teste-2026",
            "role": role,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_owner_can_create_list_update_and_deactivate_users(client: TestClient):
    provision(client, "equipe-a", "owner@example.com")
    owner_token = login(client, "equipe-a", "owner@example.com")
    user = create_user(client, owner_token, "contadora@example.com", "accountant")

    assert user["role"] == "accountant"
    assert user["is_active"] is True
    assert "password_hash" not in user

    listed = client.get("/api/v1/users", headers={"Authorization": f"Bearer {owner_token}"})
    assert listed.status_code == 200
    assert {item["email"] for item in listed.json()} == {
        "owner@example.com",
        "contadora@example.com",
    }

    updated = client.patch(
        f"/api/v1/users/{user['id']}",
        headers={"Authorization": f"Bearer {owner_token}"},
        json={"role": "assistant"},
    )
    assert updated.status_code == 200
    assert updated.json()["role"] == "assistant"

    deactivated = client.patch(
        f"/api/v1/users/{user['id']}",
        headers={"Authorization": f"Bearer {owner_token}"},
        json={"is_active": False},
    )
    assert deactivated.status_code == 200
    assert deactivated.json()["is_active"] is False

    denied_login = client.post(
        "/api/v1/auth/token",
        json={
            "tenant_slug": "equipe-a",
            "email": "contadora@example.com",
            "password": "Senha-Forte-Teste-2026",
        },
    )
    assert denied_login.status_code == 401


def test_non_manager_cannot_manage_users_and_viewer_is_read_only(client: TestClient):
    provision(client, "permissoes", "owner@example.com")
    owner_token = login(client, "permissoes", "owner@example.com")
    accountant = create_user(client, owner_token, "contabil@example.com", "accountant")
    viewer = create_user(client, owner_token, "consulta@example.com", "viewer")

    accountant_token = login(client, "permissoes", "contabil@example.com")
    headers = {"Authorization": f"Bearer {accountant_token}"}
    assert client.get("/api/v1/users", headers=headers).status_code == 403
    assert client.post(
        "/api/v1/users",
        headers=headers,
        json={
            "email": "outro@example.com",
            "full_name": "Outro Usuário",
            "password": "Senha-Forte-Teste-2026",
            "role": "accountant",
        },
    ).status_code == 403
    assert client.post(
        "/api/v1/clients",
        headers=headers,
        json={"legal_name": "Cliente Cadastrado Pela Contabilidade"},
    ).status_code == 201

    viewer_token = login(client, "permissoes", "consulta@example.com")
    viewer_headers = {"Authorization": f"Bearer {viewer_token}"}
    assert client.get("/api/v1/clients", headers=viewer_headers).status_code == 200
    assert client.post(
        "/api/v1/clients",
        headers=viewer_headers,
        json={"legal_name": "Cadastro Não Permitido"},
    ).status_code == 403
    assert accountant["role"] == "accountant"
    assert viewer["role"] == "viewer"


def test_admin_cannot_create_admin_change_owner_or_disable_self(client: TestClient):
    owner = provision(client, "escritorio-admin", "owner@example.com")
    owner_token = login(client, "escritorio-admin", "owner@example.com")
    admin = create_user(client, owner_token, "admin@example.com", "admin")
    admin_token = login(client, "escritorio-admin", "admin@example.com")
    admin_headers = {"Authorization": f"Bearer {admin_token}"}

    rejected_admin = client.post(
        "/api/v1/users",
        headers=admin_headers,
        json={
            "email": "admin2@example.com",
            "full_name": "Outro Admin",
            "password": "Senha-Forte-Teste-2026",
            "role": "admin",
        },
    )
    assert rejected_admin.status_code == 403

    edit_owner = client.patch(
        f"/api/v1/users/{owner['owner']['id']}",
        headers=admin_headers,
        json={"is_active": False},
    )
    assert edit_owner.status_code == 403

    disable_self = client.patch(
        f"/api/v1/users/{admin['id']}",
        headers=admin_headers,
        json={"is_active": False},
    )
    assert disable_self.status_code == 403

    created_staff = create_user(client, admin_token, "assistente@example.com", "assistant")
    assert created_staff["role"] == "assistant"


def test_user_management_is_scoped_to_authenticated_tenant(client: TestClient):
    provision(client, "tenant-a", "owner-a@example.com")
    token_a = login(client, "tenant-a", "owner-a@example.com")
    foreign_user = create_user(client, token_a, "colaborador-a@example.com", "accountant")

    provision(client, "tenant-b", "owner-b@example.com")
    token_b = login(client, "tenant-b", "owner-b@example.com")
    headers_b = {"Authorization": f"Bearer {token_b}"}

    listing = client.get("/api/v1/users", headers=headers_b)
    assert listing.status_code == 200
    assert {row["email"] for row in listing.json()} == {"owner-b@example.com"}

    update_foreign = client.patch(
        f"/api/v1/users/{foreign_user['id']}",
        headers=headers_b,
        json={"is_active": False},
    )
    assert update_foreign.status_code == 404



def test_client_import_preview_requires_confirmation_and_imports_atomically(client: TestClient):
    provision(client, "importacao", "owner@example.com")
    token = login(client, "importacao", "owner@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    csv_content = (
        "Razão Social;Nome Fantasia;CNPJ;E-mail;Telefone;Status\n"
        "Empresa Importada LTDA;Empresa Importada;12.345.678/0001-99;"
        "contato@importada.com.br;(93) 99999-0000;Ativo\n"
    )
    preview = client.post(
        "/api/v1/imports/clients/preview",
        headers=headers,
        files={"file": ("clientes.csv", csv_content.encode("utf-8"), "text/csv")},
    )
    assert preview.status_code == 201, preview.text
    batch = preview.json()
    assert batch["status"] == "ready"
    assert batch["total_rows"] == 1
    assert batch["rows"][0]["legal_name"] == "Empresa Importada LTDA"

    # Preview is non-destructive until the user makes an explicit confirmation.
    assert client.get("/api/v1/clients", headers=headers).json() == []

    confirmed = client.post(
        f"/api/v1/imports/clients/{batch['id']}/confirm",
        headers=headers,
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["status"] == "confirmed"
    assert confirmed.json()["imported_count"] == 1
    clients = client.get("/api/v1/clients", headers=headers).json()
    assert len(clients) == 1
    assert clients[0]["legal_name"] == "Empresa Importada LTDA"

    repeated = client.post(
        f"/api/v1/imports/clients/{batch['id']}/confirm",
        headers=headers,
    )
    assert repeated.status_code == 409


def test_invalid_client_import_is_blocked_and_reports_rows(client: TestClient):
    provision(client, "importacao-invalida", "owner@example.com")
    token = login(client, "importacao-invalida", "owner@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    csv_content = (
        "Razão Social;CNPJ;E-mail\n"
        "Empresa Válida LTDA;11.111.111/0001-11;email-invalido\n"
        "Empresa Duplicada LTDA;11.111.111/0001-11;outro@example.com\n"
    )
    preview = client.post(
        "/api/v1/imports/clients/preview",
        headers=headers,
        files={"file": ("clientes.csv", csv_content.encode("utf-8"), "text/csv")},
    )
    assert preview.status_code == 201, preview.text
    batch = preview.json()
    assert batch["status"] == "invalid"
    assert any("E-mail inválido" in error for error in batch["rows"][0]["errors"])
    assert any("duplicado dentro do arquivo" in error for error in batch["rows"][1]["errors"])

    confirmed = client.post(
        f"/api/v1/imports/clients/{batch['id']}/confirm",
        headers=headers,
    )
    assert confirmed.status_code == 409
    assert client.get("/api/v1/clients", headers=headers).json() == []


def test_client_import_supports_xlsx(client: TestClient):
    from io import BytesIO

    from openpyxl import Workbook

    provision(client, "importacao-xlsx", "owner@example.com")
    token = login(client, "importacao-xlsx", "owner@example.com")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Razão Social", "Nome Fantasia", "CNPJ", "E-mail"])
    sheet.append(["Planilha Exemplo LTDA", "Planilha Exemplo", "22.222.222/0001-22", "xlsx@example.com"])
    content = BytesIO()
    workbook.save(content)

    response = client.post(
        "/api/v1/imports/clients/preview",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "clientes.xlsx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201, response.text
    assert response.json()["status"] == "ready"
    assert response.json()["rows"][0]["email"] == "xlsx@example.com"


def test_client_import_is_scoped_to_tenant_and_rejects_viewer(client: TestClient):
    provision(client, "importacao-a", "owner-a@example.com")
    token_a = login(client, "importacao-a", "owner-a@example.com")
    uploaded = client.post(
        "/api/v1/imports/clients/preview",
        headers={"Authorization": f"Bearer {token_a}"},
        files={
            "file": (
                "clientes.csv",
                b"Razao Social;CNPJ\nEmpresa do A;33.333.333/0001-33\n",
                "text/csv",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    batch_id = uploaded.json()["id"]

    provision(client, "importacao-b", "owner-b@example.com")
    token_b = login(client, "importacao-b", "owner-b@example.com")
    headers_b = {"Authorization": f"Bearer {token_b}"}
    assert client.get(f"/api/v1/imports/clients/{batch_id}", headers=headers_b).status_code == 404
    assert client.post(
        f"/api/v1/imports/clients/{batch_id}/confirm", headers=headers_b
    ).status_code == 404

    owner_token = login(client, "importacao-a", "owner-a@example.com")
    create_user(client, owner_token, "consulta@example.com", "viewer")
    viewer_token = login(client, "importacao-a", "consulta@example.com")
    denied = client.post(
        "/api/v1/imports/clients/preview",
        headers={"Authorization": f"Bearer {viewer_token}"},
        files={"file": ("clientes.csv", b"Razao Social\nEmpresa\n", "text/csv")},
    )
    assert denied.status_code == 403


def test_client_import_rejects_unsupported_or_empty_files(client: TestClient):
    provision(client, "importacao-formato", "owner@example.com")
    token = login(client, "importacao-formato", "owner@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    unsupported = client.post(
        "/api/v1/imports/clients/preview",
        headers=headers,
        files={"file": ("clientes.xls", b"not xlsx", "application/octet-stream")},
    )
    assert unsupported.status_code == 422

    empty_rows = client.post(
        "/api/v1/imports/clients/preview",
        headers=headers,
        files={"file": ("clientes.csv", b"Razao Social;CNPJ\n", "text/csv")},
    )
    assert empty_rows.status_code == 422



def test_pending_client_import_preview_can_be_discarded(client: TestClient):
    provision(client, "importacao-descartar", "owner@example.com")
    token = login(client, "importacao-descartar", "owner@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    preview = client.post(
        "/api/v1/imports/clients/preview",
        headers=headers,
        files={
            "file": (
                "clientes.csv",
                b"Razao Social;CNPJ\nEmpresa Pendente;44.444.444/0001-44\n",
                "text/csv",
            )
        },
    )
    assert preview.status_code == 201, preview.text
    batch_id = preview.json()["id"]

    discarded = client.delete(f"/api/v1/imports/clients/{batch_id}", headers=headers)
    assert discarded.status_code == 204
    retained = client.get(f"/api/v1/imports/clients/{batch_id}", headers=headers)
    assert retained.status_code == 200
    assert retained.json()["status"] == "discarded"
    assert retained.json()["rows"] == []
    history = client.get("/api/v1/imports/clients", headers=headers)
    assert history.status_code == 200
    assert history.json()[0]["status"] == "discarded"
    assert client.get("/api/v1/clients", headers=headers).json() == []



def test_client_import_templates_require_auth_and_have_expected_headers(client: TestClient):
    unauthenticated_csv = client.get("/api/v1/imports/client-template.csv")
    unauthenticated_xlsx = client.get("/api/v1/imports/client-template.xlsx")
    assert unauthenticated_csv.status_code == 401
    assert unauthenticated_xlsx.status_code == 401

    provision(client, "modelo-planilha", "owner@example.com")
    token = login(client, "modelo-planilha", "owner@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    csv_template = client.get("/api/v1/imports/client-template.csv", headers=headers)
    assert csv_template.status_code == 200
    assert "text/csv" in csv_template.headers["content-type"]
    assert "attachment;" in csv_template.headers["content-disposition"]
    csv_text = csv_template.content.decode("utf-8-sig")
    csv_headers = csv_text.strip().splitlines()[0].split(";")
    assert csv_headers == ["Razão Social", "Nome Fantasia", "CNPJ/CPF", "E-mail", "Telefone", "Status"]

    from io import BytesIO

    from openpyxl import load_workbook

    xlsx_template = client.get("/api/v1/imports/client-template.xlsx", headers=headers)
    assert xlsx_template.status_code == 200
    workbook = load_workbook(BytesIO(xlsx_template.content), data_only=True)
    assert workbook.sheetnames == ["Clientes", "Instruções"]
    sheet = workbook["Clientes"]
    assert [cell.value for cell in sheet[1]] == csv_headers
    assert sheet["C2"].number_format == "@"
    assert "preencha" in str(workbook["Instruções"]["A2"].value).lower()
    workbook.close()


def test_import_history_tracks_discarded_and_confirmed_batches_per_tenant(client: TestClient):
    provision(client, "historico-a", "owner-a@example.com")
    token_a = login(client, "historico-a", "owner-a@example.com")
    headers_a = {"Authorization": f"Bearer {token_a}"}

    discarded_preview = client.post(
        "/api/v1/imports/clients/preview",
        headers=headers_a,
        files={
            "file": (
                "clientes-pendentes.csv",
                b"Razao Social;CNPJ\nEmpresa Descartada;55.555.555/0001-55\n",
                "text/csv",
            )
        },
    )
    assert discarded_preview.status_code == 201, discarded_preview.text
    discarded_id = discarded_preview.json()["id"]

    before_discard = client.get("/api/v1/imports/clients", headers=headers_a)
    assert before_discard.status_code == 200
    assert len(before_discard.json()) == 1
    assert before_discard.json()[0]["status"] == "ready"
    assert before_discard.json()[0]["created_by_name"] == "Responsável Contábil"

    discarded = client.delete(f"/api/v1/imports/clients/{discarded_id}", headers=headers_a)
    assert discarded.status_code == 204

    confirmed_preview = client.post(
        "/api/v1/imports/clients/preview",
        headers=headers_a,
        files={
            "file": (
                "clientes-confirmados.csv",
                b"Razao Social;CNPJ\nEmpresa Confirmada;66.666.666/0001-66\n",
                "text/csv",
            )
        },
    )
    assert confirmed_preview.status_code == 201, confirmed_preview.text
    confirmed_id = confirmed_preview.json()["id"]
    confirmed = client.post(
        f"/api/v1/imports/clients/{confirmed_id}/confirm",
        headers=headers_a,
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["rows"] == []

    forbidden_delete_confirmed = client.delete(
        f"/api/v1/imports/clients/{confirmed_id}", headers=headers_a
    )
    assert forbidden_delete_confirmed.status_code == 409

    history = client.get("/api/v1/imports/clients?limit=20", headers=headers_a)
    assert history.status_code == 200
    rows = {row["id"]: row for row in history.json()}
    assert rows[discarded_id]["status"] == "discarded"
    assert rows[discarded_id]["imported_count"] == 0
    assert rows[confirmed_id]["status"] == "confirmed"
    assert rows[confirmed_id]["imported_count"] == 1
    assert rows[confirmed_id]["confirmed_at"] is not None

    provision(client, "historico-b", "owner-b@example.com")
    token_b = login(client, "historico-b", "owner-b@example.com")
    history_b = client.get(
        "/api/v1/imports/clients",
        headers={"Authorization": f"Bearer {token_b}"},
    )
    assert history_b.status_code == 200
    assert history_b.json() == []


def test_viewer_cannot_download_import_templates_or_read_import_history(client: TestClient):
    provision(client, "historico-viewer", "owner@example.com")
    owner_token = login(client, "historico-viewer", "owner@example.com")
    create_user(client, owner_token, "consulta@example.com", "viewer")
    viewer_token = login(client, "historico-viewer", "consulta@example.com")
    headers = {"Authorization": f"Bearer {viewer_token}"}
    assert client.get("/api/v1/imports/clients", headers=headers).status_code == 403
    assert client.get("/api/v1/imports/client-template.csv", headers=headers).status_code == 403
    assert client.get("/api/v1/imports/client-template.xlsx", headers=headers).status_code == 403



def create_client_record(client: TestClient, token: str, legal_name: str) -> dict:
    response = client.post(
        "/api/v1/clients",
        headers={"Authorization": f"Bearer {token}"},
        json={"legal_name": legal_name},
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_client_task_lifecycle_and_filters(client: TestClient):
    provision(client, "tarefas", "owner@example.com")
    token = login(client, "tarefas", "owner@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    company = create_client_record(client, token, "Cliente de Teste LTDA")
    teammate = create_user(client, token, "assistente@example.com", "assistant", "Assistente Contábil")

    created = client.post(
        "/api/v1/tasks",
        headers=headers,
        json={
            "client_id": company["id"],
            "title": "Solicitar extratos bancários",
            "description": "Enviar extratos referentes ao mês anterior.",
            "kind": "document",
            "priority": "high",
            "due_date": "2026-11-05",
            "assigned_to_user_id": teammate["id"],
        },
    )
    assert created.status_code == 201, created.text
    task = created.json()
    assert task["client_name"] == "Cliente de Teste LTDA"
    assert task["assigned_to_name"] == "Assistente Contábil"
    assert task["status"] == "open"
    assert task["due_date"] == "2026-11-05"

    listed = client.get("/api/v1/tasks", headers=headers)
    assert listed.status_code == 200
    assert [item["id"] for item in listed.json()] == [task["id"]]
    filtered = client.get("/api/v1/tasks?status=open&client_id=" + company["id"], headers=headers)
    assert filtered.status_code == 200
    assert len(filtered.json()) == 1
    assert client.get("/api/v1/tasks?status=not-a-status", headers=headers).status_code == 422

    received = client.patch(
        f"/api/v1/tasks/{task['id']}",
        headers=headers,
        json={"status": "received"},
    )
    assert received.status_code == 200, received.text
    assert received.json()["status"] == "received"
    assert received.json()["completed_at"] is None

    completed = client.patch(
        f"/api/v1/tasks/{task['id']}",
        headers=headers,
        json={"status": "completed"},
    )
    assert completed.status_code == 200, completed.text
    assert completed.json()["status"] == "completed"
    assert completed.json()["completed_at"] is not None

    reopened = client.patch(
        f"/api/v1/tasks/{task['id']}",
        headers=headers,
        json={"status": "open", "due_date": None, "assigned_to_user_id": None},
    )
    assert reopened.status_code == 200, reopened.text
    assert reopened.json()["status"] == "open"
    assert reopened.json()["due_date"] is None
    assert reopened.json()["assigned_to_user_id"] is None
    assert reopened.json()["completed_at"] is None


def test_client_tasks_are_isolated_by_tenant_and_validate_client_and_assignee(client: TestClient):
    provision(client, "tarefas-a", "owner-a@example.com")
    token_a = login(client, "tarefas-a", "owner-a@example.com")
    company_a = create_client_record(client, token_a, "Cliente Particular A")
    task_response = client.post(
        "/api/v1/tasks",
        headers={"Authorization": f"Bearer {token_a}"},
        json={"client_id": company_a["id"], "title": "Enviar documentos do mês"},
    )
    assert task_response.status_code == 201, task_response.text
    task_id = task_response.json()["id"]

    provision(client, "tarefas-b", "owner-b@example.com")
    token_b = login(client, "tarefas-b", "owner-b@example.com")
    headers_b = {"Authorization": f"Bearer {token_b}"}
    assert client.get("/api/v1/tasks", headers=headers_b).json() == []
    assert client.patch(
        f"/api/v1/tasks/{task_id}", headers=headers_b, json={"status": "completed"}
    ).status_code == 404
    assert client.post(
        "/api/v1/tasks",
        headers=headers_b,
        json={"client_id": company_a["id"], "title": "Tentar vincular cliente alheio"},
    ).status_code == 404

    foreign_staff = create_user(client, token_b, "assistente-b@example.com", "assistant")
    assert client.post(
        "/api/v1/tasks",
        headers={"Authorization": f"Bearer {token_a}"},
        json={
            "client_id": company_a["id"],
            "title": "Responsável de outro escritório",
            "assigned_to_user_id": foreign_staff["id"],
        },
    ).status_code == 422


def test_viewer_can_read_but_cannot_create_or_update_client_tasks(client: TestClient):
    provision(client, "tarefas-viewer", "owner@example.com")
    owner_token = login(client, "tarefas-viewer", "owner@example.com")
    company = create_client_record(client, owner_token, "Cliente Consulta")
    task_response = client.post(
        "/api/v1/tasks",
        headers={"Authorization": f"Bearer {owner_token}"},
        json={"client_id": company["id"], "title": "Contrato social atualizado", "kind": "document"},
    )
    assert task_response.status_code == 201, task_response.text
    task_id = task_response.json()["id"]

    create_user(client, owner_token, "consulta@example.com", "viewer")
    viewer_token = login(client, "tarefas-viewer", "consulta@example.com")
    viewer_headers = {"Authorization": f"Bearer {viewer_token}"}
    assert len(client.get("/api/v1/tasks", headers=viewer_headers).json()) == 1
    assert client.post(
        "/api/v1/tasks",
        headers=viewer_headers,
        json={"client_id": company["id"], "title": "Não pode criar pendência"},
    ).status_code == 403
    assert client.patch(
        f"/api/v1/tasks/{task_id}",
        headers=viewer_headers,
        json={"status": "completed"},
    ).status_code == 403


def test_client_task_rejects_invalid_values_and_missing_title(client: TestClient):
    provision(client, "tarefas-validacao", "owner@example.com")
    token = login(client, "tarefas-validacao", "owner@example.com")
    company = create_client_record(client, token, "Cliente para Validação")
    headers = {"Authorization": f"Bearer {token}"}

    assert client.post(
        "/api/v1/tasks",
        headers=headers,
        json={"client_id": company["id"], "title": "x"},
    ).status_code == 422
    assert client.post(
        "/api/v1/tasks",
        headers=headers,
        json={"client_id": company["id"], "title": "Prioridade desconhecida", "priority": "critical"},
    ).status_code == 422
    task = client.post(
        "/api/v1/tasks",
        headers=headers,
        json={"client_id": company["id"], "title": "Pendência válida"},
    )
    assert task.status_code == 201
    assert client.patch(
        f"/api/v1/tasks/{task.json()['id']}",
        headers=headers,
        json={"title": None},
    ).status_code == 422
